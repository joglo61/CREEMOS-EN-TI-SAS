from __future__ import annotations
import os
import shutil
import tempfile
from datetime import datetime, date
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request, status, Query, BackgroundTasks
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from app.database.database import get_db
from app.core.config import settings
from app.schemas.common import SuccessResponse
from app.security.auth import get_current_user
from app.models.usuario import Usuario
from app.models.archivo import Archivo
from app.models.log import Log
from app.models.cliente import Cliente
from app.models.prestamo import Prestamo
from app.models.pago import Pago
from app.models.factura import Factura
from app.utils.excel_utils import detectar_encabezados, validar_columnas

router = APIRouter(prefix="/api/v1/excel", tags=["Excel"])


def _validar_columnas_flexible(filepath: str, tipo: str) -> list[str]:
    import openpyxl
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active
    if not ws:
        wb.close()
        raise ValueError("El archivo está vacío.")
    header_row, headers = detectar_encabezados(ws, tipo)
    wb.close()
    return validar_columnas(headers, tipo)


def _respaldar_si_existe(tipo: str, db: Session, usuario_id: int) -> None:
    activo = db.query(Archivo).filter(Archivo.tipo == tipo, Archivo.activo == True).first()
    if activo and os.path.exists(activo.ruta):
        backup_dir = os.path.join(settings.BACKUP_DIR, "excel")
        os.makedirs(backup_dir, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_name = f"{tipo}_v{activo.version}_{ts}.xlsx"
        backup_path = os.path.join(backup_dir, backup_name)
        shutil.copy2(activo.ruta, backup_path)
        from app.models.backup import Backup
        b = Backup(tipo=f"excel_{tipo}", archivo=backup_name, ruta=backup_path,
                   tamano=os.path.getsize(backup_path), usuario_id=usuario_id)
        db.add(b)
        activo.activo = False
        db.flush()


@router.post("/upload")
async def upload_excel(
    request: Request,
    background_tasks: BackgroundTasks,
    tipo: str = Query(..., pattern=r"^(clientes|financiero|cartera_creemos)$"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
):
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Solo archivos .xlsx o .xls.")
    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024
    content = await file.read()
    if len(content) > max_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Archivo demasiado grande. Máximo {settings.MAX_UPLOAD_MB} MB.")
    if len(content) < 4:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Archivo vacío o inválido.")
    magic = content[:4]
    is_xlsx = magic == b"PK\x03\x04"
    is_xls = magic[:2] == b"\xd0\xcf"
    if not (is_xlsx or is_xls):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="El archivo no parece ser un Excel válido.")

    if tipo == "cartera_creemos":
        from pathlib import Path
        project_root = Path(__file__).resolve().parent.parent.parent.parent
        target_name = "Creemos.xlsx"
        target_path = project_root / target_name
        backup_path = project_root / "backend" / "data" / f"backup_{target_name}"
        os.makedirs(project_root / "backend" / "data", exist_ok=True)
        if target_path.exists():
            import shutil
            shutil.copy2(str(target_path), str(backup_path))
        with open(str(target_path), "wb") as f:
            f.write(content)
        log = Log(
            usuario_id=current_user.id, accion="CARGAR_CARTERA_EXCEL", modulo="Excel",
            descripcion=f"Archivo cartera {target_name} actualizado.",
            direccion_ip=request.client.host if request.client else None,
        )
        db.add(log)
        db.commit()

        def _sync_cartera_background():
            import logging
            logger = logging.getLogger("app.main")
            try:
                from app.cartera.sincronizador import SincronizadorCartera
                sync = SincronizadorCartera()
                resultado = sync.ejecutar()
                logger.info(f"Sync cartera post-upload: {resultado}")
                from app.database.database import SessionLocal as MainSession
                main_db = MainSession()
                try:
                    from app.cartera.importar_a_sistema import importar_cartera_a_sistema
                    imp = importar_cartera_a_sistema(main_db, usuario_id=current_user.id)
                    logger.info(f"Import post-upload: {imp}")
                finally:
                    main_db.close()
            except Exception as e:
                logger.error(f"Sync/import cartera post-upload falló: {e}")

        background_tasks.add_task(_sync_cartera_background)
        return SuccessResponse(message=f"Cartera {target_name} actualizada. Sincronizando en segundo plano…")

    os.makedirs(settings.DATA_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    basename = os.path.basename(file.filename)
    safe_name = f"{tipo}_{ts}_{basename}"
    filepath = os.path.join(settings.DATA_DIR, safe_name)
    with open(filepath, "wb") as f:
        f.write(content)
    try:
        headers = _validar_columnas_flexible(filepath, tipo)
    except ValueError as e:
        os.remove(filepath)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception as e:
        os.remove(filepath)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Error al leer el archivo: {str(e)}")
    _respaldar_si_existe(tipo, db, current_user.id)
    ultimo = db.query(Archivo).filter(Archivo.tipo == tipo).order_by(Archivo.version.desc()).first()
    version = (ultimo.version + 1) if ultimo else 1
    archivo = Archivo(
        tipo=tipo, nombre_original=file.filename, nombre_interno=safe_name,
        ruta=filepath, version=version, usuario_id=current_user.id, activo=True,
    )
    db.add(archivo)
    log = Log(
        usuario_id=current_user.id, accion="CARGAR_EXCEL", modulo="Excel",
        descripcion=f"Archivo {tipo} ({file.filename}) v{version} cargado. Columnas: {len(headers)}.",
        direccion_ip=request.client.host if request.client else None,
    )
    db.add(log)
    db.commit()
    return SuccessResponse(message=f"Archivo {tipo} cargado correctamente (v{version}).", data={
        "tipo": tipo, "version": version, "nombre": file.filename, "columnas": headers,
    })


@router.get("/versiones")
def listar_versiones(
    tipo: str = Query("", pattern=r"^(clientes|financiero)?$"),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    q = db.query(Archivo).order_by(Archivo.id.desc())
    if tipo:
        q = q.filter(Archivo.tipo == tipo)
    items = q.limit(100).all()
    result = []
    for a in items:
        result.append({
            "id": a.id, "tipo": a.tipo, "nombre": a.nombre_original,
            "version": a.version, "activo": a.activo,
            "fecha": str(a.created_at) if a.created_at else None,
            "usuario": a.usuario.nombre if a.usuario else "",
        })
    return SuccessResponse(data={"items": result})


@router.post("/restaurar/{archivo_id}")
def restaurar_version(
    request: Request,
    archivo_id: int,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
):
    archivo = db.query(Archivo).filter(Archivo.id == archivo_id).first()
    if not archivo:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado.")
    if not os.path.exists(archivo.ruta):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="El archivo físico no existe.")

    _respaldar_si_existe(archivo.tipo, db, current_user.id)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    restored_name = f"{archivo.tipo}_restaurado_{ts}_{archivo.nombre_original}"
    restored_path = os.path.join(settings.DATA_DIR, restored_name)
    shutil.copy2(archivo.ruta, restored_path)

    ultimo = db.query(Archivo).filter(Archivo.tipo == archivo.tipo).order_by(Archivo.version.desc()).first()
    new_version = (ultimo.version + 1) if ultimo else 1
    nuevo = Archivo(
        tipo=archivo.tipo, nombre_original=f"RESTAURADO_{archivo.nombre_original}",
        nombre_interno=restored_name, ruta=restored_path,
        version=new_version, usuario_id=current_user.id, activo=True,
    )
    db.add(nuevo)
    log = Log(
        usuario_id=current_user.id, accion="RESTAURAR_EXCEL", modulo="Excel",
        descripcion=f"Versión v{archivo.version} de {archivo.tipo} restaurada como v{new_version}.",
        direccion_ip=request.client.host if request.client else None,
    )
    db.add(log)
    db.commit()
    return SuccessResponse(message=f"Versión v{archivo.version} restaurada como v{new_version}.")


@router.delete("/activo/{tipo}")
def eliminar_activo(
    request: Request,
    tipo: str,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
):
    if tipo == "cartera_creemos":
        from pathlib import Path
        project_root = Path(__file__).resolve().parent.parent.parent.parent
        fname = "Creemos.xlsx"
        fp = project_root / fname
        if not fp.exists():
            raise HTTPException(status_code=404, detail=f"Archivo {fname} no encontrado.")
        try:
            os.remove(str(fp))
        except PermissionError:
            raise HTTPException(status_code=400, detail=f"El archivo {fname} está abierto en otro programa. Ciérralo e inténtalo de nuevo.")
        log = Log(
            usuario_id=current_user.id, accion="ELIMINAR_EXCEL_CARTERA", modulo="Excel",
            descripcion=f"Archivo cartera {fname} eliminado.",
            direccion_ip=request.client.host if request.client else None,
        )
        db.add(log)
        db.commit()
        return SuccessResponse(message=f"Archivo {fname} eliminado.")
    if tipo not in ("clientes", "financiero"):
        raise HTTPException(status_code=400, detail="Tipo inválido.")
    activo = db.query(Archivo).filter(Archivo.tipo == tipo, Archivo.activo == True).first()
    if not activo:
        raise HTTPException(status_code=404, detail="No hay archivo activo para este tipo.")
    activo.activo = False
    log = Log(
        usuario_id=current_user.id, accion="ELIMINAR_EXCEL_ACTIVO", modulo="Excel",
        descripcion=f"Archivo activo {tipo} (v{activo.version}) eliminado.",
        direccion_ip=request.client.host if request.client else None,
    )
    db.add(log)
    db.commit()
    return SuccessResponse(message=f"Archivo {tipo} activo eliminado.")


@router.get("/estado")
def estado_excel(
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    from pathlib import Path
    project_root = Path(__file__).resolve().parent.parent.parent.parent
    result = {}
    for tipo in ("clientes", "financiero"):
        activo = db.query(Archivo).filter(Archivo.tipo == tipo, Archivo.activo == True).first()
        total_versiones = db.query(Archivo).filter(Archivo.tipo == tipo).count()
        if activo:
            size = os.path.getsize(activo.ruta) if os.path.exists(activo.ruta) else 0
            result[tipo] = {
                "activo": True, "nombre": activo.nombre_original,
                "version": activo.version, "tamano": size,
                "fecha": str(activo.created_at) if activo.created_at else None,
                "total_versiones": total_versiones,
            }
        else:
            result[tipo] = {"activo": False, "total_versiones": total_versiones}
    for ctipo, fname in (("cartera_creemos", "Creemos.xlsx"),):
        fp = project_root / fname
        result[ctipo] = {
            "activo": fp.exists(),
            "nombre": fname,
            "tamano": os.path.getsize(str(fp)) if fp.exists() else 0,
            "fecha": str(datetime.fromtimestamp(os.path.getmtime(str(fp)))) if fp.exists() else None,
        }
    return SuccessResponse(data=result)


@router.get("/descargar-cartera")
def descargar_cartera(_: Usuario = Depends(get_current_user)):
    from pathlib import Path
    fp = Path(__file__).resolve().parent.parent.parent.parent / "Creemos.xlsx"
    if not fp.exists():
        raise HTTPException(status_code=404, detail="Creemos.xlsx no encontrado.")
    return FileResponse(str(fp), filename="Creemos.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/exportar-todo")
def exportar_todo(
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    import openpyxl
    import traceback
    try:
        wb = openpyxl.Workbook()

        # Clientes
        ws_cli = wb.active
        ws_cli.title = "Clientes"
        ws_cli.append(["ID", "Nombre", "Cedula", "Placa", "Telefono", "Direccion", "Correo", "Estado", "Observaciones", "Creado"])
        for c in db.query(Cliente).order_by(Cliente.id).all():
            ws_cli.append([c.id, c.nombre, c.cedula, c.placa, c.telefono, c.direccion, c.correo, c.estado, c.observaciones, str(c.created_at) if c.created_at else ""])

        # Prestamos
        ws_pre = wb.create_sheet("Prestamos")
        ws_pre.append(["ID", "Cliente ID", "Cliente", "Capital Inicial", "Saldo Actual", "Valor Cuota", "Tasa Interes", "Fecha Inicio", "Fecha Proximo Pago", "Estado", "Creado"])
        for p in db.query(Prestamo).order_by(Prestamo.id).all():
            cli = db.query(Cliente).filter(Cliente.id == p.cliente_id).first()
            ws_pre.append([p.id, p.cliente_id, cli.nombre if cli else "", str(p.capital_inicial), str(p.saldo_actual), str(p.valor_cuota), str(p.tasa_interes), str(p.fecha_inicio), str(p.fecha_proximo_pago), p.estado, str(p.created_at) if p.created_at else ""])

        # Pagos
        ws_pag = wb.create_sheet("Pagos")
        ws_pag.append(["ID", "Prestamo ID", "Cliente", "Numero Factura", "Fecha Pago", "Valor Pagado", "Intereses", "Capital", "Saldo Anterior", "Saldo Nuevo", "Creado"])
        for p in db.query(Pago).order_by(Pago.id).all():
            pr = db.query(Prestamo).filter(Prestamo.id == p.prestamo_id).first()
            cli_n = ""
            if pr:
                c = db.query(Cliente).filter(Cliente.id == pr.cliente_id).first()
                if c: cli_n = c.nombre
            ws_pag.append([p.id, p.prestamo_id, cli_n, p.numero_factura, str(p.fecha_pago), str(p.valor_pagado), str(p.intereses), str(p.capital), str(p.saldo_anterior), str(p.saldo_nuevo), str(p.created_at) if p.created_at else ""])

        # Facturas
        ws_fac = wb.create_sheet("Facturas")
        ws_fac.append(["ID", "Prestamo ID", "Cliente", "Numero Factura", "Total", "Creado"])
        for f in db.query(Factura).order_by(Factura.id).all():
            f_prestamo_id = f.pago.prestamo_id if f.pago else None
            pr = db.query(Prestamo).filter(Prestamo.id == f_prestamo_id).first()
            cli_n = ""
            if pr:
                c = db.query(Cliente).filter(Cliente.id == pr.cliente_id).first()
                if c: cli_n = c.nombre
            total_val = str(f.pago.valor_pagado) if f.pago else ""
            ws_fac.append([f.id, f_prestamo_id, cli_n, f.numero_factura, total_val, str(f.created_at) if f.created_at else ""])

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        wb.save(tmp.name)
        wb.close()
        ts = date.today().isoformat()
        background_tasks.add_task(os.unlink, tmp.name)
        return FileResponse(tmp.name, filename=f"respaldo_completo_{ts}.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        tb = traceback.format_exc()
        with open(os.path.join(settings.LOGS_DIR, "export_error.log"), "a") as f:
            f.write(f"Export error: {e}\n{tb}\n")
        raise HTTPException(status_code=500, detail=f"Error al exportar: {e}")


@router.post("/guardar")
def guardar_excel(
    request: Request,
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(get_current_user),
):
    from app.utils.excel_export import exportar_a_excel
    from app.models.archivo import Archivo
    activo = db.query(Archivo).filter(Archivo.tipo == "clientes", Archivo.activo == True).first()
    if not activo:
        raise HTTPException(status_code=404, detail="No hay archivo Excel activo.")
    try:
        exportar_a_excel()
        # Also copy to the root project folder with the original name
        root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        project_root = os.path.dirname(os.path.dirname(root_dir))
        dest_name = "LISTADO CLIENTES CREEMOS EN TI SAS.xlsx"
        dest_path = os.path.join(project_root, dest_name)
        shutil.copy2(activo.ruta, dest_path)
        log = Log(
            usuario_id=current_user.id, accion="GUARDAR_EXCEL", modulo="Excel",
            descripcion=f"Datos guardados en {dest_name}.",
            direccion_ip=request.client.host if request.client else None,
        )
        db.add(log)
        db.commit()
        return SuccessResponse(message=f"Datos guardados en {dest_name} correctamente.")
    except Exception as e:
        import traceback
        log_path = os.path.join(settings.LOGS_DIR, "guardar_excel_error.log")
        with open(log_path, "a") as f:
            f.write(f"Error guardando Excel: {e}\n{traceback.format_exc()}\n")
        raise HTTPException(status_code=500, detail=f"Error al guardar en Excel: {e}")
