from __future__ import annotations
import os
from datetime import date, datetime
from decimal import Decimal
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload
from app.database.database import get_db
from app.schemas.common import SuccessResponse
from app.security.auth import get_current_user
from app.models.usuario import Usuario
from app.models.cliente import Cliente
from app.models.prestamo import Prestamo
from app.models.pago import Pago
from app.models.factura import Factura
from app.models.log import Log
from app.services.cartera_mensual import cartera_del_mes

router = APIRouter(prefix="/api/v1/reportes", tags=["Reportes"])


@router.get("/clientes")
def reporte_clientes(
    estado: str = Query("", max_length=20),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    q = db.query(Cliente).order_by(Cliente.nombre)
    if estado:
        q = q.filter(Cliente.estado == estado)
    items = q.all()
    return SuccessResponse(data={
        "items": [{
            "id": c.id, "nombre": c.nombre, "cedula": c.cedula, "placa": c.placa,
            "telefono": c.telefono, "estado": c.estado, "created_at": str(c.created_at) if c.created_at else None,
        } for c in items]
    })


@router.get("/pagos")
def reporte_pagos(
    desde: str = Query(""), hasta: str = Query(""),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    q = db.query(Pago).options(joinedload(Pago.prestamo).joinedload(Prestamo.cliente)).order_by(Pago.fecha_pago.desc())
    if desde:
        q = q.filter(Pago.fecha_pago >= date.fromisoformat(desde))
    if hasta:
        q = q.filter(Pago.fecha_pago <= date.fromisoformat(hasta))
    items = q.limit(500).all()
    return SuccessResponse(data={
        "items": [{
            "id": p.id, "numero_factura": p.numero_factura, "fecha_pago": str(p.fecha_pago),
            "valor_pagado": str(p.valor_pagado), "intereses": str(p.intereses),
            "capital": str(p.capital), "saldo_anterior": str(p.saldo_anterior),
            "saldo_nuevo": str(p.saldo_nuevo),
            "cliente": p.prestamo.cliente.nombre if p.prestamo and p.prestamo.cliente else "",
        } for p in items]
    })


@router.get("/ingresos")
def reporte_ingresos(
    desde: str = Query(""), hasta: str = Query(""),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    q = db.query(
        Pago.fecha_pago,
        (Pago.intereses + Pago.intereses_mora).label("valor")
    ).order_by(Pago.fecha_pago)
    if desde:
        q = q.filter(Pago.fecha_pago >= date.fromisoformat(desde))
    if hasta:
        q = q.filter(Pago.fecha_pago <= date.fromisoformat(hasta))
    items = q.all()
    total = sum(float(r.valor) for r in items)
    return SuccessResponse(data={
        "items": [{"fecha": str(r.fecha_pago), "valor": str(r.valor)} for r in items],
        "total": round(total),
    })


@router.get("/mora")
def reporte_mora(
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    from app.services.prestamo_service import PrestamoService
    PrestamoService(db).recalcular_estados()
    prestamos = db.query(Prestamo).options(joinedload(Prestamo.cliente)).filter(Prestamo.estado == "MORA").order_by(Prestamo.fecha_proximo_pago).all()
    return SuccessResponse(data={
        "items": [{
            "id": p.id, "cliente": p.cliente.nombre if p.cliente else "", "saldo": str(p.saldo_actual),
            "cuota": str(p.valor_cuota), "proximo_pago": str(p.fecha_proximo_pago),
            "dias_mora": (date.today() - p.fecha_proximo_pago).days if p.fecha_proximo_pago else 0,
        } for p in prestamos]
    })


@router.get("/facturas")
def reporte_facturas(
    desde: str = Query(""), hasta: str = Query(""),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    q = db.query(Factura).options(joinedload(Factura.cliente), joinedload(Factura.pago)).order_by(Factura.fecha.desc())
    if desde:
        q = q.filter(Factura.fecha >= date.fromisoformat(desde))
    if hasta:
        q = q.filter(Factura.fecha <= date.fromisoformat(hasta))
    items = q.limit(500).all()
    return SuccessResponse(data={
        "items": [{
            "id": f.id, "numero": f.numero_factura, "fecha": str(f.fecha),
            "cliente": f.cliente.nombre if f.cliente else "", "valor": str(f.pago.valor_pagado) if f.pago else "",
            "estado": f.estado,
        } for f in items]
    })


@router.get("/flujo-mensual")
def flujo_mensual(
    meses: int = Query(12, ge=1, le=60),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    from calendar import monthrange
    hoy = date.today()
    labels: list[str] = []
    ingresos: list[float] = []
    intereses: list[float] = []
    capital_recuperado: list[float] = []
    nuevos_prestamos: list[float] = []
    saldo_final: list[float] = []

    half = meses // 2
    for i in range(half, -half - 1, -1):
        m = hoy.month - i
        y = hoy.year
        while m < 1:
            m += 12
            y -= 1
        while m > 12:
            m -= 12
            y += 1
        _, last_day = monthrange(y, m)
        month_start = date(y, m, 1)
        month_end = date(y, m, last_day)
        labels.append(f"{y}-{m:02d}")

        # Ingresos (pagos) del mes
        row = db.query(
            func.coalesce(func.sum(Pago.valor_pagado), 0),
            func.coalesce(func.sum(Pago.intereses), 0),
            func.coalesce(func.sum(Pago.capital), 0),
        ).filter(Pago.fecha_pago >= month_start, Pago.fecha_pago <= month_end).first()
        tot, inte, cap = (float(r) if r else 0 for r in row) if row else (0, 0, 0)
        ingresos.append(tot)
        intereses.append(inte)
        capital_recuperado.append(cap)

        # Nuevos préstamos del mes
        new_row = db.query(func.coalesce(func.sum(Prestamo.capital_inicial), 0)).filter(
            Prestamo.fecha_inicio >= month_start, Prestamo.fecha_inicio <= month_end,
        ).first()
        nuevos_prestamos.append(float(new_row[0]) if new_row else 0)

        # Saldo de la cartera al cierre de ESE mes (histórico, desde los pagos)
        saldo_final.append(float(cartera_del_mes(db, y, m)["totales"]["saldo_final"]))

    flujo_neto = [i - n for i, n in zip(intereses, nuevos_prestamos)]

    return SuccessResponse(data={
        "labels": labels,
        "ingresos": intereses,
        "intereses": intereses,
        "capital_recuperado": capital_recuperado,
        "nuevos_prestamos": nuevos_prestamos,
        "flujo_neto": flujo_neto,
        "saldo_final": saldo_final,
    })


def _mes_valido(anio: int | None, mes: int | None) -> tuple[int, int]:
    hoy = date.today()
    return (anio or hoy.year, mes or hoy.month)


@router.get("/cartera-mensual")
def reporte_cartera_mensual(
    anio: int | None = Query(None, ge=2015, le=2100),
    mes: int | None = Query(None, ge=1, le=12),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    """Cartera por cobrar del mes (reemplazo del bloque CXCOBRAR del Excel)."""
    return SuccessResponse(data=cartera_del_mes(db, *_mes_valido(anio, mes)))


@router.get("/cartera-mensual/excel")
def exportar_cartera_mensual(
    background_tasks: BackgroundTasks,
    anio: int | None = Query(None, ge=2015, le=2100),
    mes: int | None = Query(None, ge=1, le=12),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    import tempfile
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    data = cartera_del_mes(db, *_mes_valido(anio, mes))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CXCOBRAR"
    ws.append([data["mes"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    cols = ["Fecha", "Placa", "Cliente", "Vr.credito", "Vr. Cuota", "Saldo anterior", "Fecha Inicial",
            "Fecha Final", "Dias", "Intereses", "Int. Mora", "Abono K", "Cuota", "Saldo Final"]
    ws.append(cols)
    for c in ws[3]:
        c.font = Font(bold=True)
    pago = PatternFill("solid", fgColor="DBEAFE")  # fila sombreada = pagó en el mes (como el color del Excel)
    for f in data["items"]:
        ws.append([f["fecha_desembolso"], f["placa"], f["cliente"], f["vr_credito"], f["vr_cuota"], f["saldo_anterior"],
                   f["fecha_inicial"], f["fecha_final"], f["dias"], f["intereses"] if f["pago_en_mes"] else None,
                   f["interes_mora"] or None, f["abono_capital"] if f["pago_en_mes"] else None,
                   f["cuota"] if f["pago_en_mes"] else None, f["saldo_final"]])
        if f["pago_en_mes"]:
            for c in ws[ws.max_row]:
                c.fill = pago
    t = data["totales"]
    ws.append(["TOTAL", None, None, None, None, t["saldo_anterior"], None, None, None,
               t["abono_intereses"] - t["interes_mora"], t["interes_mora"], t["abono_capital"], t["recaudo"], t["saldo_final"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    for etiqueta, valor in (("Recaudo", t["recaudo"]), ("Abono a Capital", t["abono_capital"]),
                            ("Abono Intereses", t["abono_intereses"]), ("2,5% sobre saldo anterior", t["interes_esperado"]),
                            ("Créditos / pagaron", f"{t['creditos']} / {t['pagaron']}")):
        ws.append([None, None, None, None, etiqueta, None, None, valor])
    for row in ws.iter_rows(min_row=4):
        for c in row:
            if isinstance(c.value, Decimal):
                c.value = int(c.value)
                c.number_format = "#,##0"
    for letra, ancho in zip("ABCDEFGHIJKLMN", (12, 10, 34, 13, 11, 14, 12, 12, 6, 12, 11, 12, 12, 14)):
        ws.column_dimensions[letra].width = ancho

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)
    background_tasks.add_task(os.unlink, tmp.name)
    nombre = f"cartera_{data['anio']}_{data['numero_mes']:02d}.xlsx"
    return FileResponse(tmp.name, filename=nombre,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/exportar-excel")
def exportar_excel(
    reporte: str = Query(..., pattern=r"^(clientes|pagos|ingresos|mora|facturas)$"),
    db: Session = Depends(get_db),
    _: Usuario = Depends(get_current_user),
):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from app.core.config import settings
    import os

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = reporte.capitalize()
    os.makedirs(settings.DATA_DIR, exist_ok=True)
    filename = f"reporte_{reporte}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(settings.DATA_DIR, filename)

    if reporte == "clientes":
        ws.append(["Nombre", "Cédula", "Placa", "Teléfono", "Estado", "Creado"])
        for c in db.query(Cliente).order_by(Cliente.nombre).all():
            ws.append([c.nombre, c.cedula, c.placa, c.telefono or "", c.estado, str(c.created_at or "")])
    elif reporte == "pagos":
        ws.append(["Factura", "Fecha", "Cliente", "Valor", "Interés", "Capital", "Saldo Anterior", "Saldo Nuevo"])
        q = db.query(Pago).options(joinedload(Pago.prestamo).joinedload(Prestamo.cliente)).order_by(Pago.fecha_pago.desc()).limit(500)
        for p in q:
            ws.append([p.numero_factura, str(p.fecha_pago), p.prestamo.cliente.nombre if p.prestamo and p.prestamo.cliente else "",
                       str(p.valor_pagado), str(p.intereses), str(p.capital), str(p.saldo_anterior), str(p.saldo_nuevo)])
    elif reporte == "ingresos":
        ws.append(["Fecha", "Intereses"])
        q = db.query(Pago.fecha_pago, (Pago.intereses + Pago.intereses_mora).label("total_intereses")).order_by(Pago.fecha_pago).all()
        for r in q:
            ws.append([str(r.fecha_pago), str(r.total_intereses)])
    elif reporte == "mora":
        ws.append(["Cliente", "Saldo", "Cuota", "Próximo Pago", "Días Mora"])
        q = db.query(Prestamo).options(joinedload(Prestamo.cliente)).filter(Prestamo.estado == "MORA").all()
        for p in q:
            ws.append([p.cliente.nombre if p.cliente else "", str(p.saldo_actual), str(p.valor_cuota),
                       str(p.fecha_proximo_pago), (date.today() - p.fecha_proximo_pago).days if p.fecha_proximo_pago else 0])
    elif reporte == "facturas":
        ws.append(["Factura", "Fecha", "Cliente", "Valor", "Estado"])
        q = db.query(Factura).options(joinedload(Factura.cliente), joinedload(Factura.pago)).order_by(Factura.fecha.desc()).limit(500)
        for f in q:
            ws.append([f.numero_factura, str(f.fecha), f.cliente.nombre if f.cliente else "",
                       str(f.pago.valor_pagado) if f.pago else "", f.estado])

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    wb.save(filepath)
    return FileResponse(filepath, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=filename)
