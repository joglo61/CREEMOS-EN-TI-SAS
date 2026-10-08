from __future__ import annotations
import os
from datetime import date
from decimal import Decimal
from fastapi import BackgroundTasks
from app.database.database import SessionLocal
from app.models.cliente import Cliente
from app.models.prestamo import Prestamo
from app.models.pago import Pago
from app.models.archivo import Archivo
from app.utils.excel_utils import detectar_encabezados, construir_mapa


def _get_activo_ruta() -> str | None:
    db = SessionLocal()
    try:
        a = db.query(Archivo).filter(Archivo.tipo == "clientes", Archivo.activo == True).first()
        if not a or not os.path.exists(a.ruta):
            return None
        return a.ruta
    finally:
        db.close()


def exportar_a_excel() -> None:
    ruta = _get_activo_ruta()
    if not ruta:
        return

    import openpyxl
    wb = openpyxl.load_workbook(ruta)
    ws = wb.active
    if not ws:
        wb.close()
        return

    header_row, headers = detectar_encabezados(ws, "clientes")

    db = SessionLocal()
    try:
        clients = db.query(Cliente).filter(Cliente.estado == "ACTIVO").order_by(Cliente.id).all()

        # Read original data from Excel by placa
        original_rows: dict[str, list] = {}
        for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            vals = list(row)
            placa_orig = str(vals[0]).strip() if vals[0] is not None else ""
            if placa_orig:
                original_rows[placa_orig] = vals

        # Preload prestamos for all active clients
        client_ids = [c.id for c in clients]
        all_prestamos = db.query(Prestamo).filter(
            Prestamo.cliente_id.in_(client_ids),
            Prestamo.estado.in_(["ACTIVO", "MORA"]),
        ).all()
        prestamo_por_cliente: dict[int, Prestamo | None] = {}
        for p in all_prestamos:
            if p.cliente_id not in prestamo_por_cliente:
                prestamo_por_cliente[p.cliente_id] = p

        # Preload last payment date and last factura per prestamo
        prestamo_ids = [p.id for p in all_prestamos]
        ultima_fecha_por_prestamo: dict[int, date | None] = {}
        ultima_factura_por_prestamo: dict[int, str | None] = {}
        if prestamo_ids:
            from sqlalchemy import func as sfunc
            rows = db.query(Pago.prestamo_id, sfunc.max(Pago.fecha_pago)).filter(
                Pago.prestamo_id.in_(prestamo_ids),
            ).group_by(Pago.prestamo_id).all()
            for pid, fmax in rows:
                ultima_fecha_por_prestamo[pid] = fmax
            # Get last factura per prestamo
            from sqlalchemy import or_
            subq = db.query(
                Pago.prestamo_id, sfunc.max(Pago.id).label("max_id")
            ).filter(Pago.prestamo_id.in_(prestamo_ids)).group_by(Pago.prestamo_id).subquery()
            last_pagos = db.query(Pago).join(subq, Pago.id == subq.c.max_id).all()
            for lp in last_pagos:
                ultima_factura_por_prestamo[lp.prestamo_id] = lp.numero_factura

        # Unmerge any merged cells in the data area, then clear values
        merged_to_remove = [m for m in ws.merged_cells.ranges if m.min_row > header_row]
        for m in merged_to_remove:
            ws.unmerge_cells(str(m))
        max_row = ws.max_row
        for row_idx in range(header_row + 1, max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                try:
                    ws.cell(row=row_idx, column=col_idx).value = None
                except AttributeError:
                    pass

        # Write data rows
        SKIP_PLACAS = {"0", "15000000", "CLIENTES TAMAYO", "CLIENTES DE JOGLO"}
        row_idx = header_row + 1
        for c in clients:
            if c.placa in SKIP_PLACAS:
                continue

            orig = original_rows.get(c.placa, [None] * ws.max_column)
            prestamo = prestamo_por_cliente.get(c.id)
            ultima_fecha = ultima_fecha_por_prestamo.get(prestamo.id) if prestamo else None

            # Col 1: PLACA
            ws.cell(row=row_idx, column=1, value=c.placa)
            # Col 2: MODELO (preserve original)
            ws.cell(row=row_idx, column=2, value=orig[1] if len(orig) > 1 else None)
            # Col 3: CTA INICIAL
            ws.cell(row=row_idx, column=3, value=int(prestamo.capital_inicial) if prestamo and prestamo.capital_inicial else (orig[2] if len(orig) > 2 else None))
            # Col 4: last factura number (numeric part) or preserve original
            num_factura = ultima_factura_por_prestamo.get(prestamo.id) if prestamo else None
            if num_factura:
                import re
                m = re.search(r"(\d+)$", num_factura)
                col4_val = int(m.group(1)) if m else (orig[3] if len(orig) > 3 else None)
            else:
                col4_val = orig[3] if len(orig) > 3 else None
            ws.cell(row=row_idx, column=4, value=col4_val)
            # Col 5: NOMBRE
            ws.cell(row=row_idx, column=5, value=c.nombre)
            # Col 6: ULTI FECHA (last pago or prox fecha)
            ws.cell(row=row_idx, column=6, value=ultima_fecha or (prestamo.fecha_proximo_pago if prestamo else (orig[5] if len(orig) > 5 else None)))
            # Col 7: SALDO
            ws.cell(row=row_idx, column=7, value=int(prestamo.saldo_actual) if prestamo and prestamo.saldo_actual is not None else (orig[6] if len(orig) > 6 else None))
            # Col 8: diferencia (preserve original formula if any)
            ws.cell(row=row_idx, column=8, value=orig[7] if len(orig) > 7 else None)
            # Col 9: CUOTA
            ws.cell(row=row_idx, column=9, value=int(prestamo.valor_cuota) if prestamo and prestamo.valor_cuota is not None else (orig[8] if len(orig) > 8 else None))
            # Col 10: CELULAR
            ws.cell(row=row_idx, column=10, value=c.telefono or (orig[9] if len(orig) > 9 else None))
            # Col 11: INT ACUMUL (preserve original)
            ws.cell(row=row_idx, column=11, value=orig[10] if len(orig) > 10 else None)
            # Col 12: PROX FECHA
            ws.cell(row=row_idx, column=12, value=prestamo.fecha_proximo_pago if prestamo else (orig[11] if len(orig) > 11 else None))
            # Col 13-15: preserve original
            for col_idx in range(13, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx, value=orig[col_idx - 1] if len(orig) >= col_idx else None)

            row_idx += 1

        wb.save(ruta)
    finally:
        wb.close()
        db.close()


def auto_export_background() -> None:
    try:
        exportar_a_excel()
    except Exception as e:
        import traceback
        logs_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "..", "logs")
        os.makedirs(logs_dir, exist_ok=True)
        with open(os.path.join(logs_dir, "auto_export_error.log"), "a") as f:
            f.write(f"Auto-export error: {e}\n{traceback.format_exc()}\n")


def trigger_auto_export(background_tasks: BackgroundTasks | None = None) -> None:
    if background_tasks:
        background_tasks.add_task(auto_export_background)
    else:
        auto_export_background()
