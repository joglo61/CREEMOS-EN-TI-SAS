from __future__ import annotations
import os
from datetime import date, datetime as dt_datetime
from decimal import Decimal
from dateutil.relativedelta import relativedelta
from sqlalchemy.orm import Session
from app.core.config import settings
from app.models.cliente import Cliente
from app.models.prestamo import Prestamo
from app.models.pago import Pago
from app.models.factura import Factura
from app.models.archivo import Archivo
from app.models.configuracion import Configuracion
from app.models.log import Log
from app.utils.excel_utils import detectar_encabezados, construir_mapa


class SyncService:
    def __init__(self, db: Session):
        self.db = db

    def _get_activo(self, tipo: str) -> tuple[str, int, list[str]] | None:
        a = self.db.query(Archivo).filter(Archivo.tipo == tipo, Archivo.activo == True).first()
        if not a or not os.path.exists(a.ruta):
            return None
        import openpyxl
        wb = openpyxl.load_workbook(a.ruta, read_only=True)
        ws = wb.active
        if not ws:
            wb.close()
            return None
        header_row, headers = detectar_encabezados(ws, tipo)
        wb.close()
        return a.ruta, header_row, headers

    def _read_data(self, ruta: str, header_row: int):
        import openpyxl
        wb = openpyxl.load_workbook(ruta, read_only=True)
        ws = wb.active
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > header_row:
                rows.append(row)
        wb.close()
        return rows

    def _parse_num(self, raw: str) -> int | None:
        raw = raw.strip().replace(".", "").replace(",", "")
        if raw and raw.replace("-", "").replace("+", "").isdigit():
            return int(raw)
        return None

    def _parse_date(self, raw: str) -> date | None:
        raw = raw.strip()
        if not raw:
            return None
        from datetime import datetime as dt
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
            try:
                return dt.strptime(raw, fmt).date()
            except ValueError:
                pass
        return None

    def sincronizar_clientes(self, usuario_id: int | None = None) -> dict:
        result = self._get_activo("clientes")
        if not result:
            raise ValueError("No hay archivo de clientes activo.")
        ruta, header_row, headers = result
        mapa = construir_mapa(headers, "clientes")
        mapa_fin = construir_mapa(headers, "financiero")
        creados = 0
        actualizados = 0
        errores = 0
        prestamos_creados = 0
        prestamos_actualizados = 0

        existing_placas = set(r[0] for r in self.db.query(Cliente.placa).all())
        seen_during_sync: dict[str, Cliente] = {}

        for row in self._read_data(ruta, header_row):
            if not any(row):
                continue
            vals = [str(v).strip() if v is not None else "" for v in row]
            nombre = vals[mapa["nombre"]] if "nombre" in mapa else ""
            placa = vals[mapa["placa"]] if "placa" in mapa else ""
            cedula = vals[mapa["cedula"]] if "cedula" in mapa else (placa or nombre)
            if not nombre:
                errores += 1
                continue
            telefono = vals[mapa.get("telefono", -1)] if "telefono" in mapa else ""
            direccion = vals[mapa.get("direccion", -1)] if "direccion" in mapa else ""
            correo = vals[mapa.get("correo", -1)] if "correo" in mapa else ""

            # Financial data
            capital_raw = vals[mapa_fin["capital"]] if "capital" in mapa_fin else ""
            saldo_raw = vals[mapa_fin["saldo"]] if "saldo" in mapa_fin else ""
            cuota_raw = vals[mapa_fin["cuota"]] if "cuota" in mapa_fin else ""
            prox_fecha_raw = vals[mapa_fin["fecha"]] if "fecha" in mapa_fin else ""

            capital = self._parse_num(capital_raw)
            saldo = self._parse_num(saldo_raw)
            cuota = self._parse_num(cuota_raw)
            prox_fecha = self._parse_date(prox_fecha_raw)

            # Also check raw cell values (dates may be stored as datetime objects)
            if not prox_fecha:
                raw_cell = row[11] if len(row) > 11 else None
                if isinstance(raw_cell, dt_datetime):
                    prox_fecha = raw_cell.date()
            ulti_raw = row[5] if len(row) > 5 else None
            ulti_fecha = None
            if isinstance(ulti_raw, dt_datetime):
                ulti_fecha = ulti_raw.date()
            elif isinstance(ulti_raw, str):
                ulti_fecha = self._parse_date(ulti_raw)

            # Check if we already added/updated this placa in this sync
            if placa and placa in seen_during_sync:
                existente = seen_during_sync[placa]
                if telefono: existente.telefono = telefono
                if direccion: existente.direccion = direccion
                if correo: existente.correo = correo
                actualizados += 1
                continue

            # Check existing in DB
            if placa and placa in existing_placas:
                existente = self.db.query(Cliente).filter(Cliente.placa == placa).first()
                if not existente:
                    existente = self.db.query(Cliente).filter(
                        (Cliente.cedula == (cedula or "")) | (Cliente.nombre == nombre)
                    ).first()
                if existente:
                    if nombre: existente.nombre = nombre
                    if placa: existente.placa = placa
                    if telefono: existente.telefono = telefono
                    if direccion: existente.direccion = direccion
                    if correo: existente.correo = correo
                    actualizados += 1
                    seen_during_sync[placa] = existente
                    if cedula:
                        existing_placas.add(cedula)
                    continue

            # Create new client
            c = Cliente(
                nombre=nombre, cedula=cedula or placa or nombre,
                placa=placa or cedula or nombre,
                telefono=telefono or None, direccion=direccion or None,
                correo=correo or None, estado="ACTIVO",
            )
            self.db.add(c)
            creados += 1
            seen_during_sync[placa] = c
            existing_placas.add(placa)

        self.db.commit()

        # Second pass: create/update loans from financial data
        for row in self._read_data(ruta, header_row):
            if not any(row):
                continue
            vals = [str(v).strip() if v is not None else "" for v in row]
            placa = vals[mapa["placa"]] if "placa" in mapa else ""
            if not placa:
                continue

            cliente = self.db.query(Cliente).filter(Cliente.placa == placa).first()
            if not cliente:
                continue

            capital_raw = vals[mapa_fin["capital"]] if "capital" in mapa_fin else ""
            saldo_raw = vals[mapa_fin["saldo"]] if "saldo" in mapa_fin else ""
            cuota_raw = vals[mapa_fin["cuota"]] if "cuota" in mapa_fin else ""

            capital = self._parse_num(capital_raw)
            saldo = self._parse_num(saldo_raw)
            cuota = self._parse_num(cuota_raw)

            if saldo is None and capital is None and cuota is None:
                continue

            raw_prox = row[11] if len(row) > 11 else None
            prox_fecha = None
            if isinstance(raw_prox, dt_datetime):
                prox_fecha = raw_prox.date()
            elif isinstance(raw_prox, str):
                prox_fecha = self._parse_date(raw_prox)

            raw_ulti = row[5] if len(row) > 5 else None
            ulti_fecha = None
            if isinstance(raw_ulti, dt_datetime):
                ulti_fecha = raw_ulti.date()

            existing_prestamo = self.db.query(Prestamo).filter(
                Prestamo.cliente_id == cliente.id, Prestamo.estado.in_(["ACTIVO", "MORA"]),
            ).first()

            fecha_inicio = ulti_fecha or date.today()
            if not prox_fecha:
                prox_fecha = fecha_inicio + relativedelta(months=1)

            if existing_prestamo:
                if saldo is not None:
                    existing_prestamo.saldo_actual = saldo
                if cuota is not None:
                    existing_prestamo.valor_cuota = cuota
                if capital is not None:
                    existing_prestamo.capital_inicial = capital
                existing_prestamo.fecha_proximo_pago = prox_fecha
                prestamos_actualizados += 1
            else:
                p = Prestamo(
                    cliente_id=cliente.id,
                    capital_inicial=capital or saldo or 0,
                    saldo_actual=saldo or capital or 0,
                    valor_cuota=cuota or 0,
                    tasa_interes=Decimal("2.5"),
                    fecha_inicio=fecha_inicio,
                    fecha_primer_pago=prox_fecha,
                    fecha_proximo_pago=prox_fecha,
                    estado="ACTIVO",
                )
                self.db.add(p)
                prestamos_creados += 1

        # Sync factura number from Excel receipt column (column 4, empty header, red 4-digit numbers)
        max_recibo = 0
        for row in self._read_data(ruta, header_row):
            if not any(row):
                continue
            raw = row[3] if len(row) > 3 else None
            if isinstance(raw, (int, float)):
                val = int(raw)
                if 1000 <= val <= 9999:
                    if val > max_recibo:
                        max_recibo = val
        if max_recibo > 0:
            config = self.db.query(Configuracion).first()
            if config and max_recibo >= config.siguiente_factura:
                config.siguiente_factura = max_recibo + 1

        self.db.commit()
        log = Log(
            usuario_id=usuario_id, accion="SINCRONIZAR_CLIENTES", modulo="Sincronizacion",
            descripcion=f"Clientes: {creados} creados, {actualizados} actualizados, {errores} errores. Prestamos: {prestamos_creados} creados, {prestamos_actualizados} actualizados.",
        )
        self.db.add(log)
        self.db.commit()
        return {"creados": creados, "actualizados": actualizados, "errores": errores, "prestamos_creados": prestamos_creados, "prestamos_actualizados": prestamos_actualizados}

    def sincronizar_financiero(self, usuario_id: int | None = None) -> dict:
        result = self._get_activo("financiero")
        if not result:
            raise ValueError("No hay archivo financiero activo.")
        ruta, header_row, headers = result
        mapa = construir_mapa(headers, "financiero")
        actualizados = 0
        errores = 0
        for row in self._read_data(ruta, header_row):
            if not any(row):
                continue
            vals = [str(v).strip() if v is not None else "" for v in row]
            try:
                cliente_id = ""
                if "cliente" in mapa:
                    cliente_id = vals[mapa["cliente"]]
                elif "nombre" in mapa:
                    cliente_id = vals[mapa["nombre"]]
                if not cliente_id:
                    errores += 1
                    continue
                cliente = self.db.query(Cliente).filter(
                    (Cliente.cedula == cliente_id) | (Cliente.placa == cliente_id) | (Cliente.nombre == cliente_id)
                ).first()
                if not cliente:
                    errores += 1
                    continue
                prestamo = self.db.query(Prestamo).filter(
                    Prestamo.cliente_id == cliente.id, Prestamo.estado.in_(["ACTIVO", "MORA"]),
                ).first()
                if not prestamo:
                    continue
                if "capital" in mapa and vals[mapa["capital"]]:
                    prestamo.capital_inicial = Decimal(vals[mapa["capital"]].replace(".", "").replace(",", ""))
                if "saldo" in mapa and vals[mapa["saldo"]]:
                    prestamo.saldo_actual = Decimal(vals[mapa["saldo"]].replace(".", "").replace(",", ""))
                if "cuota" in mapa and vals[mapa["cuota"]]:
                    prestamo.valor_cuota = Decimal(vals[mapa["cuota"]].replace(".", "").replace(",", ""))
                actualizados += 1
            except Exception:
                errores += 1
        self.db.commit()
        log = Log(
            usuario_id=usuario_id, accion="SINCRONIZAR_FINANCIERO", modulo="Sincronizacion",
            descripcion=f"Financiero sincronizado: {actualizados} actualizados, {errores} errores.",
        )
        self.db.add(log)
        self.db.commit()
        return {"actualizados": actualizados, "errores": errores}

    def exportar_pagos_a_excel(self, fecha_desde: date | None = None) -> str:
        import openpyxl
        export_dir = os.path.join(settings.DATA_DIR, "export")
        os.makedirs(export_dir, exist_ok=True)
        filepath = os.path.join(export_dir, f"pagos_{date.today().isoformat()}.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pagos"
        ws.append(["Factura", "Cliente", "Cedula", "Fecha", "Valor", "Interes", "Capital", "Saldo Anterior", "Saldo Nuevo", "Usuario"])
        q = self.db.query(Pago).order_by(Pago.id.desc())
        if fecha_desde:
            q = q.filter(Pago.created_at >= fecha_desde)
        for p in q.limit(1000).all():
            prestamo = self.db.query(Prestamo).filter(Prestamo.id == p.prestamo_id).first()
            cliente_nombre = ""
            cliente_cedula = ""
            if prestamo:
                c = self.db.query(Cliente).filter(Cliente.id == prestamo.cliente_id).first()
                if c:
                    cliente_nombre = c.nombre
                    cliente_cedula = c.cedula
            ws.append([
                p.numero_factura, cliente_nombre, cliente_cedula, str(p.fecha_pago),
                str(p.valor_pagado), str(p.intereses), str(p.capital),
                str(p.saldo_anterior), str(p.saldo_nuevo),
                p.usuario.nombre if p.usuario else "",
            ])
        wb.save(filepath)
        return filepath
