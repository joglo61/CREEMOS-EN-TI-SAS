import re
import openpyxl
from datetime import datetime
from pathlib import Path

EXCLUIDAS = {"COMPRAS", "CUPO", "CAJA", "CR TAMAYO", "CR SUFI", "CR FALABELLA",
             "AUX.JOGLO", "AUX. KARIME", "AUX. FCIA", "AUX.ESTELLA",
             "AUX.BCO FALABELLA", "CR CAMIONETA", "CR LIBRE INVERSION"}


def _to_int(val):
    if val is None:
        return None
    try:
        return int(round(float(str(val).replace("$", "").replace(",", ""))))
    except (ValueError, TypeError):
        return None


def parse_sheet_name(sheet_name):
    """Extract placa and suffix from sheet name."""
    name = sheet_name.strip()
    sname = name.upper()
    if sname.startswith("PRENDA "):
        sname = sname[7:].strip()

    match = re.match(r"^([A-Z0-9]+)-(\d+)$", sname)
    if match:
        return match.group(1), f"-{match.group(2)}"

    match2 = re.search(r"-(\d+)$", sname)
    if match2 and not sname.replace("-", "").replace(" ", "").isnumeric():
        base = sname[:match2.start()]
        return base.strip(), f"-{match2.group(1)}"

    return sname, ""


def parse_client_sheet(ws, sheet_name):
    result = {
        "nombre_cliente": None,
        "telefono": None,
        "valor_credito": None,
        "valor_cuota": None,
        "fecha_desembolso": None,
        "plazo_meses": None,
        "placa": None,
        "pagos": [],
        "errores": [],
    }

    for r in range(1, 8):
        row_vals = [ws.cell(r, c).value for c in range(1, 13)]
        for cell in row_vals:
            if cell is None or not isinstance(cell, str):
                continue
            cul = cell.upper().strip()
            if ("CREDITO" in cul or "PRESTAMO" in cul) and not result["nombre_cliente"]:
                result["nombre_cliente"] = cell.replace("CREDITO", "").replace("PRESTAMO", "").strip()
            if "TELEFONO" in cul or "TEL." in cul or "CEL" in cul:
                idx = row_vals.index(cell)
                if idx + 1 < len(row_vals) and row_vals[idx + 1] is not None:
                    result["telefono"] = str(row_vals[idx + 1])
            if "VALOR CREDITO" in cul:
                idx = row_vals.index(cell)
                for off in [1, 2, 3]:
                    v = row_vals[idx + off] if idx + off < len(row_vals) else None
                    if v is not None:
                        try:
                            result["valor_credito"] = _to_int(v)
                            if result["valor_credito"]:
                                break
                        except (ValueError, TypeError):
                            pass
            if "CUOTA" in cul and "$" in cell:
                try:
                    result["valor_cuota"] = _to_int(cell.split("$")[1])
                except (ValueError, IndexError):
                    pass
            if "FECHA DESEMBOLSO" in cul:
                for token in cell.replace("/", "-").split():
                    for fmt in ("%d-%m-%Y", "%B %d/%Y", "%d-%B-%Y"):
                        try:
                            result["fecha_desembolso"] = datetime.strptime(token, fmt)
                        except ValueError:
                            pass
            if "PLAZO" in cul:
                m = re.search(r"(\d+)", cell)
                if m:
                    result["plazo_meses"] = int(m.group(1))
            if "PLACA" in cul:
                idx = row_vals.index(cell)
                for off in [1, 2, 3]:
                    v = row_vals[idx + off] if idx + off < len(row_vals) else None
                    if v is not None:
                        result["placa"] = str(v).strip()
                        break

    if not result["placa"]:
        base, _ = parse_sheet_name(sheet_name)
        result["placa"] = base

    data_start = None
    for r in range(1, min(20, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, 12)]
        if vals and vals[0] == "PAGOS":
            data_start = r + 1
            break

    if data_start is None:
        for r in range(1, min(20, ws.max_row + 1)):
            v = ws.cell(r, 1).value
            if v == 1 or v == "1":
                data_start = r
                break

    if data_start is None:
        result["errores"].append("No se encontró el inicio de la tabla de pagos")
        return result

    ultimo_np = 0
    for r in range(data_start, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 11)]
        # Fila fantasma (plantilla con fórmulas en 0): no tiene cuota (F),
        # ni fecha de pago + recibo (D+B). No es un pago real.
        es_pago = (vals[5] is not None) or (vals[3] is not None and vals[1] is not None)
        if not es_pago:
            continue
        try:
            np = int(float(str(vals[0]))) if vals[0] is not None else None
        except (ValueError, TypeError):
            np = None
        if np is None:
            # Columna A con fórmula sin caché: la numeración es secuencial
            np = ultimo_np + 1
        ultimo_np = np

        result["pagos"].append({
            "numero_pago": np,
            "recibo": str(vals[1]) if vals[1] is not None else None,
            "fecha_ultimo_pago": vals[2].date() if isinstance(vals[2], datetime) else None,
            "fecha_pago": vals[3].date() if isinstance(vals[3], datetime) else None,
            "dias": _to_int(vals[4]),
            "cuota": _to_int(vals[5]),
            "interes": _to_int(vals[6]),
            "saldo_intereses": _to_int(vals[7]),
            "capital": _to_int(vals[8]),
            "saldo_real": _to_int(vals[9]),
        })

    return result


MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
         "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

def _es_mes(texto):
    if not texto or not isinstance(texto, str):
        return False
    t = texto.upper().strip()
    for m in MESES:
        if t.startswith(m):
            return True
    return False


class JogloParser:
    def __init__(self, filepath):
        self.filepath = Path(filepath)
        self.creditos = []
        self.pagos = []
        self.snapshots = []
        self.errores = []

    def parse_all(self):
        wb = openpyxl.load_workbook(str(self.filepath), data_only=True)

        for sheet_name in wb.sheetnames:
            if sheet_name in EXCLUIDAS:
                continue
            if sheet_name == "CXCOBRAR":
                self._parse_cxcobrar(wb)
                continue

            ws = wb[sheet_name]
            result = parse_client_sheet(ws, sheet_name)
            if result["errores"]:
                self.errores.append(f"Hoja '{sheet_name}': {'; '.join(result['errores'])}")
            if not result["pagos"] and not result["valor_credito"]:
                continue

            base_placa, sufijo = parse_sheet_name(sheet_name)

            self.creditos.append({
                "placa": base_placa,
                "sufijo_credito": sufijo,
                "nombre_cliente": result["nombre_cliente"] or sheet_name,
                "telefono": result["telefono"],
                "valor_credito": result["valor_credito"],
                "valor_cuota": result["valor_cuota"],
                "fecha_desembolso": result["fecha_desembolso"],
                "plazo_meses": result["plazo_meses"],
                "prenda": None,
                "tiene_historial_detallado": True,
            })

            for p in result["pagos"]:
                p["placa"] = base_placa
                p["sufijo"] = sufijo
                self.pagos.append(p)

        wb.close()
        return self.creditos, self.pagos, self.snapshots

    def _find_last_cxcobrar_block(self, ws):
        last = None
        r = 1
        while r <= ws.max_row:
            v = ws.cell(r, 1).value
            es_informe = v and isinstance(v, str) and "INFORME RECAUDO" in v.upper()
            es_mes = _es_mes(v)
            if es_informe or es_mes:
                if es_informe:
                    month_label = ws.cell(r + 1, 1).value
                    if month_label:
                        month_label = str(month_label).strip()
                    header_row = r + 3
                else:
                    month_label = str(v).strip()
                    header_row = r + 2

                last = (month_label, header_row)

                dr = header_row + 1
                while dr <= ws.max_row:
                    dv = ws.cell(dr, 1).value
                    if dv is None:
                        dr += 1
                        continue
                    try:
                        int(float(str(dv)))
                    except (ValueError, TypeError):
                        break
                    dr += 1
                r = dr
            else:
                r += 1
        return last

    def _parse_cxcobrar(self, wb):
        ws = wb["CXCOBRAR"]
        block = self._find_last_cxcobrar_block(ws)
        if not block:
            return

        month_label, header_row = block
        dr = header_row + 1

        while dr <= ws.max_row:
            dv = ws.cell(dr, 1).value
            if dv is None:
                dr += 1
                continue
            try:
                int(float(str(dv)))
            except (ValueError, TypeError):
                break

            row = {
                "placa": str(ws.cell(dr, 3).value or "").strip(),
                "cliente": str(ws.cell(dr, 5).value or "").strip(),
                "fecha_original": ws.cell(dr, 2).value.date() if isinstance(ws.cell(dr, 2).value, datetime) else ws.cell(dr, 2).value,
                "vr_credito": _to_int(ws.cell(dr, 6).value),
                "vr_cuota": _to_int(ws.cell(dr, 7).value),
                "saldo_anterior": _to_int(ws.cell(dr, 8).value),
                "fecha_inicial": ws.cell(dr, 9).value.date() if isinstance(ws.cell(dr, 9).value, datetime) else ws.cell(dr, 9).value,
                "fecha_final": ws.cell(dr, 10).value.date() if isinstance(ws.cell(dr, 10).value, datetime) else ws.cell(dr, 10).value,
                "dias": _to_int(ws.cell(dr, 11).value),
                "intereses": _to_int(ws.cell(dr, 12).value),
                "interes_mora": _to_int(ws.cell(dr, 13).value),
                "abono_k": _to_int(ws.cell(dr, 14).value),
                "cuota": _to_int(ws.cell(dr, 15).value),
                "saldo_final": _to_int(ws.cell(dr, 16).value),
                "prenda": str(ws.cell(dr, 17).value or "").strip(),
                "mes_reportado": month_label,
            }
            self.snapshots.append(row)
            dr += 1
