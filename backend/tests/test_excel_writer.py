"""Tests para app.cartera.excel_writer — creación de bloques y escritura de pagos."""
import openpyxl
import pytest
from datetime import date, datetime

from app.cartera.excel_writer import (
    localizar_bloques, crear_bloque_mes, escribir_pagos,
    mes_label, parse_mes_label, color_mes,
)


def _make_workbook(path, ultimo_mes="JULIO DE 2026"):
    """Crea un xlsx sintético con un bloque mensual en CXCOBRAR y una hoja individual."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CXCOBRAR"

    ws.cell(1, 1).value = ultimo_mes
    ws.cell(2, 12).value = " "
    ws.cell(2, 15).value = " "
    headers = [None, "Fecha ", "Placa ", "Cliente", "Vr.credito", "Vr. Cuota",
               "Saldo anterior", "Fecha Inicial ", "Fecha Final ", "Dias",
               "Intereses ", "Int. Mora", "Abono K", "Cuota ", "Saldo Final "]
    for c, h in enumerate(headers, start=1):
        if h:
            ws.cell(3, c).value = h

    # Fila 4: crédito vigente con pago
    ws.cell(4, 1).value = 1
    ws.cell(4, 2).value = datetime(2022, 4, 5)
    ws.cell(4, 3).value = "SNY460"
    ws.cell(4, 4).value = "CLIENTE UNO"
    ws.cell(4, 5).value = 36000000
    ws.cell(4, 6).value = 1165000
    ws.cell(4, 7).value = 30000000
    ws.cell(4, 8).value = datetime(2026, 6, 2)
    ws.cell(4, 9).value = datetime(2026, 7, 2)
    ws.cell(4, 10).value = "=+I4-H4"
    ws.cell(4, 11).value = "=(G4*0.025)/30*J4"
    ws.cell(4, 13).value = "=+N4-K4-L4"
    ws.cell(4, 14).value = 1165000
    ws.cell(4, 15).value = "=+G4-M4"

    # Fila 5: crédito SALDADO (saldo final 0) → no debe arrastrarse
    ws.cell(5, 1).value = "=1+A4"
    ws.cell(5, 2).value = datetime(2021, 1, 10)
    ws.cell(5, 3).value = "AAA111"
    ws.cell(5, 4).value = "CLIENTE SALDADO"
    ws.cell(5, 5).value = 10000000
    ws.cell(5, 6).value = 500000
    ws.cell(5, 7).value = 100000
    ws.cell(5, 8).value = datetime(2026, 7, 1)
    ws.cell(5, 9).value = datetime(2026, 7, 1)
    ws.cell(5, 10).value = "=+I5-H5"
    ws.cell(5, 11).value = "=(G5*0.025)/30*J5"
    ws.cell(5, 13).value = "=+N5-K5-L5"
    ws.cell(5, 14).value = 500000
    ws.cell(5, 15).value = 0  # saldado

    # Totales + resumen
    ws.cell(6, 7).value = "=SUM(G4:G5)"
    ws.cell(6, 11).value = "=SUM(K4:K5)"
    ws.cell(6, 12).value = "=SUM(L4:L5)"
    ws.cell(6, 13).value = "=SUM(M4:M5)"
    ws.cell(6, 14).value = "=SUM(N4:N5)"
    ws.cell(6, 15).value = "=SUM(O4:O5)"
    ws.cell(8, 5).value = f"Recaudo de {ultimo_mes.split()[0].capitalize()} de 2026"
    ws.cell(8, 8).value = "=+N6"
    ws.cell(9, 5).value = "Abono a Capital "
    ws.cell(10, 5).value = "Abono Intereses "

    # Hoja individual para SNY460
    sh = wb.create_sheet("SNY460")
    sh.cell(1, 3).value = "CREDITO  CLIENTE UNO"
    sh.cell(2, 3).value = "VALOR CREDITO"
    sh.cell(2, 4).value = 36000000
    sh.cell(5, 3).value = "PLACA  SNY460"
    sh.cell(6, 1).value = "PAGOS"
    sh.cell(7, 3).value = datetime(2022, 4, 5)
    sh.cell(7, 4).value = datetime(2022, 4, 5)
    sh.cell(7, 10).value = "=+D2"
    sh.cell(8, 1).value = 1
    sh.cell(8, 2).value = "REC-001"
    sh.cell(8, 3).value = "=+C7"
    sh.cell(8, 4).value = datetime(2026, 7, 2)
    sh.cell(8, 5).value = 30
    sh.cell(8, 6).value = 1165000
    sh.cell(8, 7).value = "=D2*2.5%/30*E8"
    sh.cell(8, 9).value = "=+F8-G8"
    sh.cell(8, 10).value = "=+J7-I8"

    wb.save(path)
    wb.close()


def test_parse_mes_label():
    assert parse_mes_label("AGOSTO DE 2026") == (2026, 8)
    assert parse_mes_label("ENERO  DE 2026") == (2026, 1)
    assert parse_mes_label("OTRA COSA") is None


def test_mes_label():
    assert mes_label(2026, 9) == "SEPTIEMBRE DE 2026"


def test_color_mes_rota():
    assert color_mes(1) != color_mes(2)
    assert color_mes(13) == color_mes(1)


def test_localizar_bloques(tmp_path):
    f = tmp_path / "test.xlsx"
    _make_workbook(str(f))
    wb = openpyxl.load_workbook(str(f))
    bloques = localizar_bloques(wb["CXCOBRAR"])
    wb.close()
    assert len(bloques) == 1
    b = bloques[0]
    assert b["label"] == "JULIO DE 2026"
    assert b["datos_inicio"] == 4
    assert b["datos_fin"] == 5
    assert b["totales"] == 6
    assert b["resumen"] == 8


def test_crear_bloque_mes(tmp_path, monkeypatch):
    import app.cartera.excel_writer as ew

    class _Agosto2026(date):
        @classmethod
        def today(cls):
            return cls(2026, 8, 15)

    monkeypatch.setattr(ew, "date", _Agosto2026)
    f = tmp_path / "test.xlsx"
    _make_workbook(str(f))

    altas = [{
        "fecha_desembolso": datetime(2026, 8, 5), "placa": "NEW001",
        "cliente": "CLIENTE NUEVO SISTEMA", "vr_credito": 5000000,
        "vr_cuota": 300000, "saldo_anterior": 5000000,
        "fecha_inicial": datetime(2026, 8, 5),
    }]
    res = crear_bloque_mes(str(f), altas)
    assert res is not None
    assert res["label"] == "AGOSTO DE 2026"  # mes siguiente al último bloque (JULIO)

    wb = openpyxl.load_workbook(str(f))
    ws = wb["CXCOBRAR"]
    bloques = localizar_bloques(ws)
    assert len(bloques) == 2
    nuevo = bloques[-1]

    # Debe tener 2 filas: SNY460 (arrastrado) + NEW001 (alta). AAA111 saldado → fuera.
    placas = []
    for r in range(nuevo["datos_inicio"], nuevo["datos_fin"] + 1):
        placas.append(ws.cell(r, 3).value)
    assert placas == ["SNY460", "NEW001"]

    # Fila arrastrada: G = saldo anterior (O anterior era fórmula → None en este archivo
    # sintético sin valores cacheados, así que cae al fallback de G)
    r1 = nuevo["datos_inicio"]
    assert ws.cell(r1, 3).value == "SNY460"
    assert ws.cell(r1, 7).value is not None  # saldo anterior poblado
    # Fórmulas presentes
    assert str(ws.cell(r1, 10).value).startswith("=+I")
    assert "0.025" in str(ws.cell(r1, 11).value)
    assert str(ws.cell(r1, 15).value).startswith("=+G")
    # Cuota N vacía hasta que pague
    assert ws.cell(r1, 14).value is None

    # Alta: H = I = fecha desembolso
    r2 = r1 + 1
    assert ws.cell(r2, 3).value == "NEW001"
    assert ws.cell(r2, 8).value == datetime(2026, 8, 5)
    assert ws.cell(r2, 9).value == datetime(2026, 8, 5)

    # Totales y resumen
    assert str(ws.cell(nuevo["totales"], 7).value).startswith("=SUM(G")
    assert "Recaudo" in ws.cell(nuevo["resumen"], 5).value
    wb.close()


def test_crear_bloque_mes_noop_si_ya_existe(tmp_path):
    f = tmp_path / "test.xlsx"
    _make_workbook(str(f), ultimo_mes=mes_label(date.today().year, date.today().month))
    res = crear_bloque_mes(str(f), [])
    assert res is None  # el mes actual ya tiene bloque


def test_escribir_pagos(tmp_path):
    f = tmp_path / "test.xlsx"
    # Bloque del mes actual para que escriba ahí
    _make_workbook(str(f), ultimo_mes=mes_label(date.today().year, date.today().month))

    pagos = [{
        "placa": "SNY460", "fecha_pago": date(2026, 8, 10), "valor_pagado": 1200000,
        "intereses_mora": 0, "numero_factura": "FACT-000777",
        "dias": 30, "intereses": 750000, "capital": 450000, "saldo_nuevo": 29550000,
    }]
    res = escribir_pagos(str(f), pagos)
    assert res["escritos_bloque"] == 1
    assert res["escritos_hoja"] == 1

    wb = openpyxl.load_workbook(str(f))
    ws = wb["CXCOBRAR"]
    bloques = localizar_bloques(ws)
    r = bloques[-1]["datos_inicio"]
    # N acumulado: 1165000 previo + 1200000
    assert ws.cell(r, 14).value == 2365000.0
    assert ws.cell(r, 9).value == datetime(2026, 8, 10)
    # Relleno aplicado
    assert ws.cell(r, 4).fill.patternType == "solid"
    # Hoja individual: nueva fila con recibo FACT-000777
    sh = wb["SNY460"]
    assert sh.cell(9, 2).value == "FACT-000777"
    assert sh.cell(9, 4).value == datetime(2026, 8, 10)
    wb.close()

    # Segunda vez: dedup — no duplica
    res2 = escribir_pagos(str(f), pagos)
    assert res2["escritos_bloque"] == 0
    assert res2["escritos_hoja"] == 0
