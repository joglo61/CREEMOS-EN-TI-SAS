"""Migración única Excel → sistema: un taxi con 2 préstamos, mora preservada, saldo cuadrado, no se repite."""
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest

from app.database.database import SessionLocal
from app.models.pago import Pago
from app.models.prestamo import Prestamo
from migrar_excel import migrar

COLS = ["Fecha", "Placa", "Cliente", "Vr.credito", "Vr. Cuota", "Saldo anterior", "Fecha Inicial",
        "Fecha Final", "Dias", "Intereses", "Int. Mora", "Abono K", "Cuota", "Saldo Final"]


def _bloque(ws, fila, mes, filas):
    ws.cell(fila, 1).value = mes
    for c, h in enumerate(COLS, start=2):
        ws.cell(fila + 2, c).value = h
    for i, f in enumerate(filas):
        for c, v in enumerate(f, start=2):
            ws.cell(fila + 3 + i, c).value = v
    return fila + 3 + len(filas) + 2  # fila en blanco = fin del bloque


def _excel(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CXCOBRAR"
    d1, d2 = datetime(2024, 1, 10), datetime(2026, 1, 20)
    # Taxi XYZ123 con DOS préstamos (mismo cliente) + otro cliente sin pago en octubre
    sep = [
        [d1, "XYZ123", "ANA PEREZ", 40_000_000, 1_500_000, 30_000_000, datetime(2026, 8, 10), datetime(2026, 9, 10), 31, 750_000, 0, 750_000, 1_500_000, 29_250_000],
        [d2, "XYZ123", "ANA PEREZ", 5_000_000, 300_000, 5_000_000, datetime(2026, 8, 20), datetime(2026, 9, 20), 31, 125_000, 0, 175_000, 300_000, 4_825_000],
        [datetime(2025, 5, 5), "ABC999", "LUIS GOMEZ", 20_000_000, 900_000, 10_000_000, datetime(2026, 8, 5), datetime(2026, 9, 5), 31, 250_000, 0, 650_000, 900_000, 9_350_000],
    ]
    oct_ = [
        # pagó con 7 días de mora: K + L + M = N
        [d1, "XYZ123", "ANA PEREZ", 40_000_000, 1_500_000, 29_250_000, datetime(2026, 9, 10), datetime(2026, 10, 17), 37, 731_250, 8_750, 760_000, 1_500_000, 28_490_000],
        [d2, "XYZ123", "ANA PEREZ", 5_000_000, 300_000, 4_825_000, datetime(2026, 9, 20), datetime(2026, 10, 20), 30, 120_625, 0, 179_375, 300_000, 4_645_625],
        [datetime(2025, 5, 5), "ABC999", "LUIS GOMEZ", 20_000_000, 900_000, 9_350_000, datetime(2026, 9, 5), datetime(2026, 9, 5), None, None, None, None, None, 9_350_000],
    ]
    nxt = _bloque(ws, 1, "SEPTIEMBRE DE 2026", sep)
    _bloque(ws, nxt, "OCTUBRE DE 2026", oct_)
    wb.save(path)


@pytest.fixture
def db():
    s = SessionLocal()
    yield s
    s.close()


def test_dry_run_cuadra_y_no_escribe(db, tmp_path):
    f = tmp_path / "c.xlsx"
    _excel(f)
    inf = migrar(db, f, dry_run=True)
    assert inf["prestamos"] == 3
    assert inf["clientes"] == 2
    assert inf["cuadra"] and inf["saldo_migrado"] == Decimal(28_490_000 + 4_645_625 + 9_350_000)
    assert db.query(Prestamo).count() == 0


def test_migracion_real(db, tmp_path):
    f = tmp_path / "c.xlsx"
    _excel(f)
    migrar(db, f, dry_run=False)

    xyz = db.query(Prestamo).filter(Prestamo.placa == "XYZ123").order_by(Prestamo.fecha_inicio).all()
    assert len(xyz) == 2  # dos préstamos del mismo taxi, no se fusionan
    assert xyz[0].cliente_id == xyz[1].cliente_id
    assert xyz[0].saldo_actual == Decimal(28_490_000)
    assert xyz[0].fecha_proximo_pago == date(2026, 11, 17)  # un mes después del último pago

    pagos = db.query(Pago).filter(Pago.prestamo_id == xyz[0].id).order_by(Pago.fecha_pago).all()
    assert [p.valor_pagado for p in pagos] == [Decimal(1_500_000), Decimal(1_500_000)]
    assert pagos[-1].intereses_mora == Decimal(8_750)  # la mora del bloque se conserva
    assert pagos[-1].saldo_nuevo == Decimal(28_490_000)

    abc = db.query(Prestamo).filter(Prestamo.placa == "ABC999").one()
    assert abc.fecha_proximo_pago == date(2026, 10, 5)
    # Estado recalculado al migrar (5 días de gracia)
    assert abc.estado == ("MORA" if (date.today() - date(2026, 10, 5)).days > 5 else "ACTIVO")
    assert db.query(Pago).filter(Pago.prestamo_id == abc.id).count() == 1  # octubre no pagó

    with pytest.raises(SystemExit):  # nunca se migra dos veces
        migrar(db, f, dry_run=False)


def test_base_limpia_borra_datos_de_prueba(db, tmp_path):
    from app.models.cliente import Cliente
    from app.models.factura import Factura
    from app.services.cliente_service import ClienteService
    from app.services.pago_service import PagoService

    prueba = ClienteService(db).crear(
        nombre="CLIENTE DE PRUEBA", cedula="123", placa="TEST01", telefono=None, direccion=None, correo=None,
        observaciones=None, capital_inicial=Decimal(1_000_000), valor_cuota=Decimal(100_000),
        fecha_inicio=date(2026, 1, 1), fecha_primer_pago=date(2026, 2, 1))["prestamo"]
    PagoService(db).registrar(prueba.id, Decimal(100_000), date(2026, 2, 1), usuario_id=1)

    f = tmp_path / "c.xlsx"
    _excel(f)
    inf = migrar(db, f, dry_run=False, base_limpia=True)
    assert inf["pagos_app_en_previos"]  # se informó el pago de prueba descartado
    assert db.query(Cliente).filter(Cliente.placa == "TEST01").count() == 0
    assert db.query(Factura).filter(Factura.numero_factura.like("FACT-%")).count() == 0
    assert db.query(Prestamo).count() == 3
